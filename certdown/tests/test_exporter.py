#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""通用导出层测试：字段映射 + 多格式 + 可扩展性。

运行：python tests/test_exporter.py
"""

from __future__ import annotations

import csv
import json
import sys
import unittest
import zipfile
from dataclasses import dataclass
from pathlib import Path
from tempfile import TemporaryDirectory

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import cert_downloader as cd            # noqa: E402
import exporter                         # noqa: E402


# --------------------------------------------------------------------------------------
# 测试数据：刻意混用 camelCase / snake_case / 中文键 / 嵌套结构
# --------------------------------------------------------------------------------------

MIXED = [
    {"name": "张三", "idNo": "A1", "meta": {"score": 91}, "tags": ["优", "快"]},
    {"姓名": "李四", "id_no": "A2", "meta": {"score": 88}, "tags": ["良"]},
    {"NAME": "王五", "id-number": "A3", "meta": {"score": 75}, "tags": []},
]


@dataclass
class Row:
    """对象型记录。"""

    index: int
    record_id: str
    name: str
    saved_as: str
    status: str


class TestPathResolution(unittest.TestCase):
    """路径解析与命名风格自适应。"""

    def test_parse_path(self) -> None:
        self.assertEqual(exporter.parse_path("a.b.c"), ["a", "b", "c"])
        self.assertEqual(exporter.parse_path("items[0].url"), ["items", 0, "url"])
        self.assertEqual(exporter.parse_path("data['姓名']"), ["data", "姓名"])

    def test_name_conversions(self) -> None:
        self.assertEqual(exporter.to_snake("createdAt"), "created_at")
        self.assertEqual(exporter.to_snake("Created-At"), "created_at")
        self.assertEqual(exporter.to_camel("created_at"), "createdAt")
        self.assertEqual(exporter.to_camel("created_at", upper_first=True), "CreatedAt")

    def test_variants(self) -> None:
        v = exporter.name_variants("createdAt")
        self.assertIn("created_at", v)
        self.assertIn("createdAt", v)
        self.assertIn("created-at", v)

    def test_nested_and_index(self) -> None:
        data = {"a": {"b": [{"c": "deep"}]}}
        self.assertEqual(exporter.get_path(data, "a.b[0].c"), "deep")
        self.assertIsNone(exporter.get_path(data, "a.b[9].c"))
        self.assertEqual(exporter.get_path(data, "a.b[0].zz", default="兜底"), "兜底")

    def test_get_path_style_fallback(self) -> None:
        """同一字段三种写法都要能取到。"""
        for rec in ({"createdAt": 1}, {"created_at": 1}, {"created-at": 1}):
            with self.subTest(rec=rec):
                self.assertEqual(exporter.get_path(rec, "createdAt"), 1)


class TestFieldMap(unittest.TestCase):
    """字段映射：别名、候选、顺序、缺失兜底。"""

    def test_multi_source_fallback(self) -> None:
        """第一个候选取不到时降级到下一个（camelCase / snake_case / 中文）。"""
        fm = exporter.FieldMap.from_config([
            {"key": "姓名", "sources": ["name", "姓名", "NAME"]},
            {"key": "编号", "sources": ["idNo", "id_no", "id-number"]},
        ])
        rows = [fm.row(r) for r in MIXED]
        print(f"\n[多候选兜底] {rows}")
        self.assertEqual([r[0] for r in rows], ["张三", "李四", "王五"])
        self.assertEqual([r[1] for r in rows], ["A1", "A2", "A3"])

    def test_nested_source(self) -> None:
        fm = exporter.FieldMap.from_config([{"key": "分数", "sources": ["meta.score"]}])
        self.assertEqual([fm.row(r)[0] for r in MIXED], [91, 88, 75])

    def test_missing_uses_default(self) -> None:
        fm = exporter.FieldMap.from_config([
            {"key": "缺失列", "sources": ["nope", "also_nope"], "default": "N/A"},
        ])
        self.assertEqual(fm.row(MIXED[0])[0], "N/A")

    def test_skip_empty(self) -> None:
        """空字符串视为没取到，继续找下一个候选。"""
        fm = exporter.FieldMap.from_config([
            {"key": "v", "sources": ["empty", "fallback"], "skip_empty": True},
            {"key": "v2", "sources": ["empty", "fallback"], "skip_empty": False},
        ])
        rec = {"empty": "", "fallback": "FB"}
        self.assertEqual(fm.row(rec)[0], "FB")       # 跳过空值
        self.assertEqual(fm.row(rec)[1], "")         # 不跳过，取到空串

    def test_column_order_and_alias(self) -> None:
        fm = exporter.FieldMap.from_config([{"key": "B"}, {"key": "A"}, {"key": "C"}])
        self.assertEqual(fm.columns, ["B", "A", "C"])
        renamed = fm.rename({"B": "第二列", "A": "第一列"}).only(["第一列", "第二列"])
        self.assertEqual(renamed.columns, ["第一列", "第二列"])

    def test_exclude(self) -> None:
        fm = exporter.FieldMap.from_config([{"key": "A"}, {"key": "B"}])
        self.assertEqual(fm.exclude(["A"]).columns, ["B"])

    def test_auto_infer_with_aliases(self) -> None:
        """自动推断列名 + 中文表头别名 + 自定义顺序。"""
        fm = exporter.FieldMap.from_records(
            MIXED[0:1], aliases={"name": "姓名", "idNo": "编号"}, order=["编号", "姓名"])
        print(f"\n[自动推断] {fm.columns}")
        self.assertEqual(fm.columns[:2], ["编号", "姓名"])

    def test_object_and_dataclass_records(self) -> None:
        """dict / dataclass / 普通对象都能取到值。"""
        fm = exporter.FieldMap.from_config([
            {"key": "姓名", "sources": ["name"]},
            {"key": "文件", "sources": ["saved_as", "savedAs"]},
        ])
        rows = [Row(1, "fr1", "张三", "a.jpg", "ok"), Row(2, "fr2", "李四", "b.jpg", "ok")]
        self.assertEqual([fm.row(r) for r in rows], [["张三", "a.jpg"], ["李四", "b.jpg"]])

    def test_formatter_builtin_and_custom(self) -> None:
        exporter.register_formatter("mask", lambda v: str(v)[:1] + "***")
        fm = exporter.FieldMap.from_config([
            {"key": "分数", "sources": ["meta.score"], "formatter": "int"},
            {"key": "脱敏", "sources": ["idNo"], "formatter": "mask"},
            {"key": "标签", "sources": ["tags"], "formatter": "join"},
        ])
        row = fm.row(MIXED[0])
        print(f"\n[转换器] {row}")
        self.assertEqual(row[0], 91)
        self.assertEqual(row[1], "A***")
        self.assertEqual(row[2], "优、快")


class TestFormats(unittest.TestCase):
    """各导出格式。"""

    def setUp(self) -> None:
        self.tmp = TemporaryDirectory()
        self.d = Path(self.tmp.name)
        self.mapping = exporter.FieldMap.from_config([
            {"key": "姓名", "sources": ["name", "姓名", "NAME"]},
            {"key": "编号", "sources": ["idNo", "id_no", "id-number"]},
            {"key": "分数", "sources": ["meta.score"]},
        ])

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def test_csv_content_and_bom(self) -> None:
        p = self.d / "r.csv"
        res = exporter.export(MIXED, p, mapping=self.mapping)
        raw = p.read_bytes()
        self.assertTrue(raw.startswith(b"\xef\xbb\xbf"), "CSV 应带 UTF-8 BOM")
        text = raw.decode("utf-8-sig")
        self.assertIn("姓名,编号,分数", text)
        self.assertEqual(res.row_count, 3)
        self.assertEqual(res.columns, ["姓名", "编号", "分数"])
        self.assertEqual(res.mime, "text/csv; charset=utf-8")
        print(f"\n[CSV]\n{text.strip()}")

    def test_csv_without_bom(self) -> None:
        p = self.d / "nobom.csv"
        exporter.export(MIXED, p, mapping=self.mapping, bom=False)
        self.assertFalse(p.read_bytes().startswith(b"\xef\xbb\xbf"))

    def test_json_orient(self) -> None:
        p = self.d / "r.json"
        exporter.export(MIXED, p, mapping=self.mapping)
        data = json.loads(p.read_text(encoding="utf-8"))
        self.assertEqual(data[0]["姓名"], "张三")
        self.assertEqual(data[2]["编号"], "A3")

        p2 = self.d / "r2.json"
        exporter.export(MIXED, p2, mapping=self.mapping, orient="columns")
        cols = json.loads(p2.read_text(encoding="utf-8"))
        self.assertEqual(cols["姓名"], ["张三", "李四", "王五"])

    def test_jsonl(self) -> None:
        p = self.d / "r.jsonl"
        exporter.export(MIXED, p, mapping=self.mapping)
        lines = [x for x in p.read_text(encoding="utf-8").splitlines() if x]
        self.assertEqual(len(lines), 3)
        self.assertEqual(json.loads(lines[0])["姓名"], "张三")

    def test_txt_delimiter(self) -> None:
        p = self.d / "r.txt"
        exporter.export(MIXED, p, mapping=self.mapping, delimiter=" | ")
        lines = p.read_text(encoding="utf-8-sig").strip().splitlines()
        print(f"\n[TXT]\n" + "\n".join(lines))
        self.assertEqual(lines[0], "姓名 | 编号 | 分数")
        self.assertEqual(lines[1], "张三 | A1 | 91")

    def test_markdown(self) -> None:
        p = self.d / "r.md"
        exporter.export(MIXED, p, mapping=self.mapping)
        text = p.read_text(encoding="utf-8")
        self.assertTrue(text.startswith("| 姓名 | 编号 | 分数 |"))
        self.assertIn("|---|---|---|", text)

    def test_tsv(self) -> None:
        p = self.d / "r.tsv"
        exporter.export(MIXED, p, mapping=self.mapping)
        self.assertIn("\t", p.read_text(encoding="utf-8-sig"))

    def test_xlsx_structure(self) -> None:
        """校验 OOXML 包结构与内容（不依赖第三方库）。"""
        p = self.d / "r.xlsx"
        exporter.export(MIXED, p, mapping=self.mapping)
        with zipfile.ZipFile(p) as zf:
            names = set(zf.namelist())
            for need in ("[Content_Types].xml", "xl/workbook.xml",
                         "xl/worksheets/sheet1.xml", "xl/styles.xml"):
                self.assertIn(need, names)
            sheet = zf.read("xl/worksheets/sheet1.xml").decode("utf-8")
        self.assertIn("张三", sheet)
        self.assertIn("<v>91</v>", sheet, "数字应写成数值单元格而不是字符串")

    def test_xlsx_readable_by_openpyxl(self) -> None:
        """用第三方库回读，证明生成的是合法 Excel 文件。"""
        try:
            import openpyxl
        except ImportError:                                # pragma: no cover
            self.skipTest("未安装 openpyxl")
        p = self.d / "r.xlsx"
        exporter.export(MIXED, p, mapping=self.mapping, sheet_name="证书报告")
        wb = openpyxl.load_workbook(p)
        ws = wb.active
        print(f"\n[XLSX] 表名={ws.title} 尺寸={ws.dimensions}")
        self.assertEqual(ws.title, "证书报告")
        self.assertEqual([c.value for c in ws[1]], ["姓名", "编号", "分数"])
        self.assertEqual([c.value for c in ws[2]], ["张三", "A1", 91])
        self.assertEqual(ws["C4"].value, 75)

    def test_extension_autofill(self) -> None:
        """没写扩展名时按格式自动补全。"""
        res = exporter.export(MIXED, self.d / "report", fmt="xlsx", mapping=self.mapping)
        self.assertTrue(res.path.endswith("report.xlsx"))
        self.assertTrue(Path(res.path).exists())

    def test_format_inference(self) -> None:
        self.assertEqual(exporter.infer_format("a.xlsx").name, "xlsx")
        self.assertEqual(exporter.infer_format("a.JSON").name, "json")
        self.assertEqual(exporter.get_format("md").mime, "text/markdown; charset=utf-8")
        self.assertEqual(exporter.get_format("xlsx").mime,
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def test_export_bytes_and_headers(self) -> None:
        data, res = exporter.export_bytes(MIXED, fmt="csv", mapping=self.mapping)
        self.assertTrue(data.startswith(b"\xef\xbb\xbf"))
        self.assertIn("张三", data.decode("utf-8-sig"))
        headers = exporter.download_headers(res, "报告.csv")
        print(f"\n[下载头] {headers}")
        self.assertEqual(headers["Content-Type"], "text/csv; charset=utf-8")
        self.assertIn("报告.csv", headers["Content-Disposition"])

    def test_auto_mapping_without_config(self) -> None:
        """不传映射时自动推断列。"""
        data, res = exporter.export_bytes([{"a": 1, "b": 2}], fmt="json")
        self.assertEqual(res.columns, ["a", "b"])


class TestExtensibility(unittest.TestCase):
    """新增格式不需要改核心逻辑。"""

    def test_register_custom_format(self) -> None:
        """注册一种「竖线分隔 + 统计尾行」的自定义格式。"""

        def writer(fp, table, spec, options):
            lines = ["|".join(map(str, table.columns))]
            lines += ["|".join("" if v is None else str(v) for v in r) for r in table.rows]
            lines.append(f"共 {len(table.rows)} 行")
            payload = ("\n".join(lines) + "\n").encode(options.get("encoding", spec.encoding))
            if options.get("bom", spec.bom):
                payload = exporter.BOMS["utf-8"] + payload
            fp.write(payload)

        exporter.register_format(exporter.FormatSpec(
            name="pipe", extensions=("pipe",), mime="text/plain; charset=utf-8",
            writer=writer, encoding="utf-8", bom=False))

        with TemporaryDirectory() as tmp:
            res = exporter.export([{"a": 1}, {"a": 2}], Path(tmp) / "x.pipe")
            text = Path(res.path).read_text(encoding="utf-8")
            print(f"\n[自定义格式]\n{text.strip()}")
            self.assertIn("a\n1\n2\n共 2 行", text)
            self.assertEqual(res.format, "pipe")
            self.assertIn("pipe", exporter.list_formats())

    def test_duplicate_register_guard(self) -> None:
        spec = exporter.get_format("csv")
        with self.assertRaises(ValueError):
            exporter.register_format(spec)


class TestIntegration(unittest.TestCase):
    """与下载器报告的集成。"""

    def setUp(self) -> None:
        self.tmp = TemporaryDirectory()
        self.d = Path(self.tmp.name)

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def _jobs(self) -> list:
        return [
            cd.Job(index=1, record_id="fr1", name="张三", id_no="A1", status="ok",
                   saved_as="张三A1.jpg",
                   attachment=cd.Attachment(obj_key="k", name="orig.jpg",
                                            ext="jpg", size=1, kind="image")),
            cd.Job(index=2, record_id="fr2", name="李四", id_no="A2", status="failed",
                   error="超时",
                   attachment=cd.Attachment(obj_key="k2", name="o2.png",
                                            ext="png", size=1, kind="image")),
        ]

    def test_default_csv_unchanged(self) -> None:
        """默认报告的表头/内容必须与旧版硬编码 CSV 一致。"""
        p = self.d / "download_report.csv"
        cd.write_report(p, self._jobs())
        rows = list(csv.reader(p.read_text(encoding="utf-8-sig").splitlines()))
        self.assertEqual(rows[0], ["序号", "记录ID", "姓名", "身份证号",
                                   "原文件名", "保存文件名", "状态", "错误信息"])
        self.assertEqual(rows[1][:4], ["1", "fr1", "张三", "A1"])
        self.assertEqual(rows[2][7], "超时")
        self.assertTrue(p.read_bytes().startswith(b"\xef\xbb\xbf"))

    def test_other_formats(self) -> None:
        for fmt, ext in (("xlsx", "xlsx"), ("json", "json"), ("txt", "txt"), ("md", "md")):
            with self.subTest(fmt=fmt):
                p = self.d / f"r.{ext}"
                res = cd.write_report(p, self._jobs(), fmt=fmt)
                self.assertTrue(p.exists() and p.stat().st_size > 0)
                self.assertEqual(res.format, fmt)

    def test_custom_mapping_preset(self) -> None:
        p = self.d / "r.csv"
        res = cd.write_report(p, self._jobs(), mapping="cert_minimal")
        self.assertEqual(res.columns, ["姓名", "身份证号", "保存文件名", "状态"])

    def test_inline_mapping_config(self) -> None:
        """直接传配置列表（新增字段不用改代码）。"""
        p = self.d / "r.csv"
        res = cd.write_report(p, self._jobs(), mapping=[
            {"key": "姓名", "sources": ["name"]},
            {"key": "是否成功", "sources": ["status"], "formatter": "bool_zh"},
        ])
        self.assertEqual(res.columns, ["姓名", "是否成功"])
        self.assertIn("是", p.read_text(encoding="utf-8-sig"))


if __name__ == "__main__":
    unittest.main(verbosity=2)
